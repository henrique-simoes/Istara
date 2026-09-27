"""D-9, R1: a report downloads as Markdown, Word or CSV, every claim with its evidence trail."""

from __future__ import annotations

import app.core.agentic  # noqa: F401  (import-order guard)

import csv
import io
import json
import uuid

import pytest
from httpx import ASGITransport, AsyncClient

from app.main import app
from app.models.database import async_session, init_db

QUOTE = "Month end is a nightmare because the bank export never matches the receipts."


async def _seed_report():
    from app.models.code_application import CodeApplication
    from app.models.document import Document, DocumentSource, DocumentStatus
    from app.models.finding import Fact, Insight, Nugget
    from app.models.project import Project
    from app.models.project_report import ProjectReport
    from app.models.research_validity import EvidenceUnit, ResearchEvidenceEdge

    await init_db()
    s = uuid.uuid4().hex[:8]
    project_id = f"proj-export-{s}"
    async with async_session() as db:
        db.add(Project(id=project_id, name="Export"))
        db.add(Document(id=f"doc-{s}", project_id=project_id, title="p02.md", file_name="p02.md",
                        status=DocumentStatus.READY, source=DocumentSource.USER_UPLOAD,
                        content_text=QUOTE))
        db.add(EvidenceUnit(id=f"eu-{s}", project_id=project_id, source_document_id=f"doc-{s}",
                            source_id=f"nugget:n-{s}", stable_id=f"n-{s}#EU-0001", unit_index=1,
                            unit_type="source_span", source_text=QUOTE,
                            source_location="p02.md#chars=120-197"))
        db.add(Nugget(id=f"n-{s}", project_id=project_id, text=QUOTE, source="p02.md"))
        db.add(ResearchEvidenceEdge(id=str(uuid.uuid4()), project_id=project_id,
                                    source_type="nugget", source_id=f"n-{s}", relation="grounded_in",
                                    target_type="evidence_unit", target_id=f"eu-{s}",
                                    evidence_unit_id=f"eu-{s}"))
        db.add(CodeApplication(id=str(uuid.uuid4()), project_id=project_id, code_id="reconciliation",
                               evidence_unit_id=f"eu-{s}", promotion_status="accepted",
                               reconciliation_status="accepted", source_text=QUOTE))
        db.add(Fact(id=f"f-{s}", project_id=project_id, text="Reconciling bank exports costs hours.",
                    nugget_ids=json.dumps([f"n-{s}"])))
        db.add(Insight(id=f"i-{s}", project_id=project_id, text="An untraced insight.",
                       fact_ids=json.dumps([])))
        report = ProjectReport(id=f"r-{s}", project_id=project_id, title="Harbor findings",
                               finding_ids_json=json.dumps([f"f-{s}", f"n-{s}", f"i-{s}"]),
                               executive_summary="Owners lose time reconciling.")
        db.add(report)
        await db.commit()
    return project_id, f"r-{s}"


@pytest.mark.asyncio
async def test_report_exports_carry_the_evidence_trail(admin_auth_headers, monkeypatch):
    from app.config import settings

    project_id, report_id = await _seed_report()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        url = f"/api/reports/{project_id}/{report_id}/export"
        md = await ac.get(f"{url}?format=md", headers=admin_auth_headers)
        table = await ac.get(f"{url}?format=csv", headers=admin_auth_headers)
        word = await ac.get(f"{url}?format=docx", headers=admin_auth_headers)
        bad = await ac.get(f"{url}?format=pdf", headers=admin_auth_headers)
        monkeypatch.setattr(settings, "team_mode", True)
        stranger = await ac.get(f"{url}?format=md")

    assert md.status_code == 200 and "attachment" in md.headers["content-disposition"]
    assert "2 of 3 findings traced to a source span" in md.text
    assert f"“{QUOTE}” — p02.md, chars=120-197 (coding: accepted)" in md.text
    assert "No evidence trail" in md.text
    rows = list(csv.DictReader(io.StringIO(table.text)))
    assert {r["finding_type"] for r in rows} == {"fact", "nugget", "insight"}
    assert any(r["document"] == "p02.md" and r["coding"] == "accepted" for r in rows)
    from docx import Document as DocxDocument

    text = "\n".join(p.text for p in DocxDocument(io.BytesIO(word.content)).paragraphs)
    assert "Harbor findings" in text and QUOTE in text
    assert bad.status_code == 400
    assert stranger.status_code in {401, 403}
